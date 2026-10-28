import base64
import csv
import io

from odoo import api, fields, models
from odoo.exceptions import UserError

ITEM_HEADERS = ('item code', 'code', 'default_code', 'sku', 'item', 'product', 'material', 'material code', 'item name', 'name')
QTY_HEADERS = ('counted', 'counted quantity', 'counted qty', 'physical', 'physical quantity', 'physical qty', 'count', 'quantity', 'qty')


class MosCountImport(models.TransientModel):
    """Fill the counted quantities of a sheet from a spreadsheet (CSV or Excel): a column for the item (its code, or its name) and one
    for the counted quantity. Nothing is entered unless every row is understood."""

    _name = 'mos.count.import'
    _description = 'Import counts from a spreadsheet'

    count_id = fields.Many2one('mos.cycle.count', required=True, readonly=True)
    file = fields.Binary(required=True)
    filename = fields.Char()

    def _rows(self):
        stored = self.file
        raw = stored.content if hasattr(stored, 'content') else base64.b64decode(stored or b'')
        name = (self.filename or '').lower()
        if name.endswith(('.xlsx', '.xlsm')) or raw[:2] == b'PK':
            try:
                from openpyxl import load_workbook

                sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
                return [list(row) for row in sheet.iter_rows(values_only=True)]
            except Exception as error:  # a damaged or unsupported workbook
                raise UserError(self.env._('The Excel file could not be read: %s', error))
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('latin-1')
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        return list(csv.reader(io.StringIO(text), dialect))

    @staticmethod
    def _find(headers, wanted):
        for number, header in enumerate(headers):
            if str(header or '').strip().lower() in wanted:
                return number
        return None

    def action_import(self):
        self.ensure_one()
        count = self.count_id
        if count.state != 'counting':
            raise UserError(self.env._('The sheet is not open for counting.'))
        rows = [row for row in self._rows() if any(cell not in (None, '') for cell in row)]
        if not rows:
            raise UserError(self.env._('The file is empty.'))
        item_col, qty_col = self._find(rows[0], ITEM_HEADERS), self._find(rows[0], QTY_HEADERS)
        if item_col is None or qty_col is None:
            raise UserError(self.env._('The first row must have a column for the item (code or name) and one for the counted quantity.'))
        by_key = {}
        for line in count.line_ids:
            for key in (line.product_id.default_code, line.product_id.name, line.product_id.display_name):
                if key:
                    by_key.setdefault(str(key).strip().lower(), line)
        problems, entries, seen = [], {}, {}
        for number, row in enumerate(rows[1:], start=2):
            key = str(row[item_col] if item_col < len(row) and row[item_col] is not None else '').strip().lower()
            if not key:
                continue
            line = by_key.get(key)
            if not line:
                problems.append(self.env._('Row %(n)s: "%(item)s" is not on this count sheet.', n=number, item=key))
                continue
            try:
                qty = float(str(row[qty_col]).replace(',', '') if qty_col < len(row) else '')
            except ValueError:
                problems.append(self.env._('Row %(n)s: "%(qty)s" is not a number.', n=number, qty=row[qty_col] if qty_col < len(row) else ''))
                continue
            if qty < 0:
                problems.append(self.env._('Row %s: a counted quantity cannot be negative.', number))
                continue
            if line.id in seen:
                problems.append(self.env._('Row %(n)s: the item is already on row %(first)s.', n=number, first=seen[line.id]))
                continue
            seen[line.id] = number
            entries[line.id] = qty
        if problems:
            raise UserError(self.env._('Nothing was imported:\n%s', '\n'.join(problems[:15]) + (self.env._('\n...and %s more.', len(problems) - 15) if len(problems) > 15 else '')))
        for line in count.line_ids.filtered(lambda l: l.id in entries):
            line.counted_qty = entries[line.id]
        count.message_post(body=self.env._('%(n)s counted quantities imported from %(file)s.', n=len(entries), file=self.filename or ''))
        return {'type': 'ir.actions.act_window_close'}


class MosCountChanges(models.TransientModel):
    """The reviewer asks for changes: what is wrong, and which lines the counter has to count again."""

    _name = 'mos.count.changes'
    _description = 'Ask for changes to a count'

    count_id = fields.Many2one('mos.cycle.count', required=True, readonly=True)
    note = fields.Text(required=True)
    line_ids = fields.Many2many('mos.cycle.count.line', string='Count again', domain="[('count_id', '=', count_id)]")

    @api.model
    def default_get(self, fields_list):
        values = super().default_get(fields_list)
        count = self.env['mos.cycle.count'].browse(values.get('count_id') or self.env.context.get('default_count_id'))
        if count:
            values['line_ids'] = [(6, 0, count.line_ids.filtered(lambda l: abs(l.variance) >= 0.05).ids)]
        return values

    def action_confirm(self):
        self.ensure_one()
        self.count_id.action_request_changes(self.note, self.line_ids)
        return {'type': 'ir.actions.act_window_close'}
